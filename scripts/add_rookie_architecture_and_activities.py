from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


SOURCE = Path('/Users/adieu/Desktop/학교/폴리텍/최종과제/과제관련/시스원_이력서_김연동_기타활동추가_복구본.docx')
TARGET = SOURCE.with_name('시스원_이력서_김연동_신입의정석아키텍처_활동추가_최종본.docx')
ARCHITECTURE = Path('/tmp/rookie_architecture.png')


def font(size, bold=False):
    candidates = [
        '/System/Library/Fonts/AppleSDGothicNeo.ttc',
        '/System/Library/Fonts/Supplemental/AppleGothic.ttf',
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size, index=1 if bold else 0)
        except OSError:
            continue
    return ImageFont.load_default()


def draw_box(draw, xy, title, lines, fill, outline, title_fill):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = xy
    draw.text((x1 + 22, y1 + 18), title, fill=title_fill, font=font(30, True))
    for index, line in enumerate(lines):
        draw.text((x1 + 22, y1 + 68 + index * 34), line, fill='#243047', font=font(22))


def arrow(draw, start, end, fill='#64748b'):
    draw.line([start, end], fill=fill, width=6)
    x1, y1 = start
    x2, y2 = end
    draw.polygon([(x2, y2), (x2 - 18, y2 - 10), (x2 - 18, y2 + 10)], fill=fill)


def create_architecture():
    image = Image.new('RGB', (1800, 1080), '#f8fafc')
    draw = ImageDraw.Draw(image)
    draw.text((900, 35), '신입의 정석 (Rookie Playbook) 시스템 아키텍처', anchor='ma', fill='#172033', font=font(40, True))
    draw.text((900, 88), 'React SPA · Spring Boot · PostgreSQL(RDS) · JWT · RAG 기반 AI 문의', anchor='ma', fill='#526176', font=font(24))

    draw_box(draw, (80, 220, 500, 490), '사용자 / 관리자', ['웹 브라우저', 'React SPA', '문서·교육·Q&A·공지'], '#e8f1ff', '#78a8ef', '#2159ae')
    draw_box(draw, (650, 190, 1150, 520), '애플리케이션 서버', ['Spring Boot REST API', 'Controller · Service · Repository', 'JWT 인증 / ROLE_USER·ADMIN', '공통 예외·응답 처리'], '#eafaf1', '#69c394', '#16764b')
    draw_box(draw, (1300, 220, 1720, 490), '데이터 계층', ['PostgreSQL (AWS RDS)', '24개 도메인 테이블', 'JPA / MyBatis', '진도·문서·질문·알림'], '#fff4e5', '#e6ae62', '#a85d08')
    draw_box(draw, (260, 700, 720, 900), 'AI 문의 응답', ['문서 임베딩', 'pgvector 벡터 검색', 'Groq LLM / Google Embedding', 'RAG 근거 기반 답변'], '#f5edff', '#b28ae3', '#6b39a8')
    draw_box(draw, (1040, 700, 1540, 900), '배포 / 운영', ['Docker 이미지', 'AWS EC2 배포', 'GitHub Actions', 'PostgreSQL RDS 연동'], '#eef2f7', '#9aa9bd', '#40536b')

    arrow(draw, (500, 355), (650, 355))
    arrow(draw, (1150, 355), (1300, 355))
    arrow(draw, (820, 520), (560, 700), '#8a5bb5')
    arrow(draw, (980, 520), (1250, 700), '#64748b')
    image.save(ARCHITECTURE, 'PNG')


def add_activity_rows(document):
    table = document.tables[5]
    for row in list(table.rows)[1:]:
        if not any(cell.text.strip() for cell in row.cells):
            row._tr.getparent().remove(row._tr)
    rows = [
        ['2026', '2026 한이음 드림업 프로젝트 참여', '한이음', 'AI 기반 소프트웨어 프로젝트 기획 및 수행'],
        ['2026', '창업경진대회 참가', '한국폴리텍대학', '프로젝트 아이디어 발표 및 사업화 가능성 검토'],
    ]
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value


def add_architecture_below_projects(document):
    target = next(p for p in document.paragraphs if p.text.strip() == '인턴 · 실무경험')
    caption = document.add_paragraph()
    caption.alignment = 1
    caption.add_run('신입의 정석 - 시스템 아키텍처').font.size = Pt(9)
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = 1
    image_paragraph.add_run().add_picture(str(ARCHITECTURE), width=Inches(6.2))
    target._p.addprevious(caption._p)
    target._p.addprevious(image_paragraph._p)


create_architecture()
copy2(SOURCE, TARGET)
document = Document(TARGET)
add_activity_rows(document)
add_architecture_below_projects(document)
document.save(TARGET)
print(TARGET)