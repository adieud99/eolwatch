from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.table import Table


BROKEN = Path('/Users/adieu/Desktop/학교/폴리텍/최종과제/과제관련/시스원_이력서_김연동_기타활동추가_최종본.docx')
ORIGINAL = Path('/Users/adieu/Desktop/학교/폴리텍/최종과제/과제관련/시스원_이력서_김연동.docx')
TARGET = BROKEN.with_name('시스원_이력서_김연동_기타활동추가_복구본.docx')


def replace_table_text(table, values):
    for row, row_values in zip(table.rows, values):
        for cell, value in zip(row.cells, row_values):
            cell.text = value


copy2(BROKEN, TARGET)
document = Document(TARGET)
original = Document(ORIGINAL)

broken_qualification = document.tables[4]._tbl
clean_qualification = deepcopy(original.tables[4]._tbl)
activity_xml = deepcopy(original.tables[3]._tbl)

broken_qualification.getparent().replace(broken_qualification, clean_qualification)

activity_heading = next(
    paragraph._p
    for paragraph in document.paragraphs
    if paragraph.text.strip() == '5. 기타활동'
)
other_heading = next(
    paragraph._p
    for paragraph in document.paragraphs
    if paragraph.text.strip() == '6. 기타사항'
)
activity_heading.getparent().remove(activity_heading)
new_heading = deepcopy(other_heading)
heading_texts = new_heading.xpath('.//w:t')
heading_texts[0].text = '5. 기타활동'
for text_node in heading_texts[1:]:
    text_node.text = ''
other_heading.addprevious(new_heading)
new_heading.addnext(activity_xml)

activity_table = Table(activity_xml, document._body)
replace_table_text(
    activity_table,
    [
        ['일자', '수상·활동명', '주관기관', '비고'],
        [
            '2026.05 ~ 2026.07',
            '기업 홈페이지 개선 및 연구데이터 분석 프로젝트 수행(팀장)',
            '(주)이에이포스',
            '고용노동부 청년일경험지원사업 프로젝트',
        ],
    ],
)

document.save(TARGET)
print(TARGET)