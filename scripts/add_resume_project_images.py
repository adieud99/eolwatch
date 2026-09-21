from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path('/Users/adieu/Desktop/학교/폴리텍/최종과제/과제관련/시스원_이력서_김연동_인프라강화_수정본.docx')
TARGET = SOURCE.with_name('시스원_이력서_김연동_사진보강_최종본.docx')
ASSETS = Path('/tmp/resume_assets')


def insert_after(paragraph):
    new_paragraph = paragraph._parent.add_paragraph()
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


def add_image_after(paragraph, image_path, caption):
    caption_paragraph = insert_after(paragraph)
    caption_paragraph.alignment = 1
    caption_run = caption_paragraph.add_run(caption)
    caption_run.font.size = Pt(8)

    image_paragraph = insert_after(caption_paragraph)
    image_paragraph.alignment = 1
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(2.25))
    return image_paragraph


copy2(SOURCE, TARGET)
document = Document(TARGET)

# The user moved the existing project images below the project tables.
# Append the replacement project-specific assets in the same area.
anchor = document.paragraphs[18]
assets = [
    (ASSETS / 'bbik' / 'arch_updated.jpg', '삑 Bbik - 서비스 아키텍처'),
    (ASSETS / 'rookie' / 'login_screen.png', '신입의 정석 - 웹 로그인 화면'),
    (ASSETS / 'stockcast' / 'odoo-stock.png', 'StockCast - Odoo 재고 화면'),
    (ASSETS / 'stockcast' / 'nfc.png', 'StockCast - NFC 입출고 화면'),
]
for image_path, caption in assets:
    anchor = add_image_after(anchor, image_path, caption)

document.save(TARGET)
print(TARGET)