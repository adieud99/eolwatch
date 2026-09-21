from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Inches, Pt


SOURCE = Path('/Users/adieu/Desktop/학교/폴리텍/최종과제/과제관련/시스원_이력서_김연동.docx')
TARGET = SOURCE.with_name('시스원_이력서_김연동_인프라강화_수정본.docx')
ASSETS = Path('/tmp/eolwatch_resume_assets')


def replace_paragraph(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.size = Pt(9)


def add_project_note(cell, text):
    paragraph = cell.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(9)


def add_images(cell, image_paths, caption):
    caption_paragraph = cell.add_paragraph()
    caption_run = caption_paragraph.add_run(caption)
    caption_run.font.size = Pt(8)
    for image_path in image_paths:
        image_paragraph = cell.add_paragraph()
        image_paragraph.alignment = 1
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(2.15))


copy2(SOURCE, TARGET)
document = Document(TARGET)

core = [
    '- 배포 자동화·CI/CD: 삑의 GitHub Actions·Fly.io 자동 배포와 신입의 정석의 Docker·AWS EC2 배포 흐름을 구축하고, 브랜치·PR·Secrets 기반 운영을 적용함.',
    '- IaC·AWS 운영: StockCast에서 Terraform으로 EC2·보안그룹·RDS·CloudWatch·예산을 코드화하고, S3 + DynamoDB 원격 상태 및 잠금을 구성함.',
    '- 장애 대응·품질: StockCast의 Docker Compose·systemd 워치독·CloudWatch 자동 복구와 pytest 93건·pip-audit·terraform validate 게이트로 운영 안정성을 검증함.',
]
for index, text in zip(range(12, 15), core):
    replace_paragraph(document.paragraphs[index], text)

project_1 = document.tables[7]
add_project_note(project_1.cell(2, 1), '인프라 포인트: GitHub Actions에서 backend 변경 감지·Secrets 관리·Fly.io 원격 빌드와 자동 배포를 구성함.')

project_2 = document.tables[8]
add_project_note(project_2.cell(2, 1), '인프라 포인트: PostgreSQL(AWS RDS) 연결 Docker 이미지를 AWS EC2에 배포 테스트하고, GitHub Actions·TeamCity CI와 브랜치·PR 흐름을 검증함.')

project_3 = document.tables[9]
project_3.cell(5, 1).paragraphs[7].text = (
    '– 운영 이슈 해결: Odoo 메모리 부족은 인스턴스 상향·스왑으로, '
    'Web NFC의 HTTP 제한은 HTTPS 적용으로, AMI 변경에 따른 서버 교체는 Terraform ignore_changes로 해결'
)
add_project_note(project_3.cell(5, 1), '인프라 포인트: Terraform·Docker Compose·Caddy·systemd·CloudWatch를 연결해 재현 가능한 배포와 무응답 프로세스 자동 복구를 구현함.')

add_images(project_1.cell(7, 1), [ASSETS / 'bbik.png'], 'GitHub 저장소 대표 이미지')
add_images(project_2.cell(7, 1), [ASSETS / 'rookie.png'], 'GitHub 저장소 대표 이미지')
add_images(project_3.cell(7, 1), [ASSETS / 'stockcast-kpi.png', ASSETS / 'stockcast-health.png'], 'GitHub 공개 화면 캡처')

document.save(TARGET)
print(TARGET)