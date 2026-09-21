from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document


SOURCE = Path('/Users/adieu/Desktop/학교/폴리텍/최종과제/과제관련/시스원_이력서_김연동_사진보강_최종본.docx')
TARGET = SOURCE.with_name('시스원_이력서_김연동_기타활동추가_최종본.docx')


def set_cell(cell, text):
    cell.text = text


copy2(SOURCE, TARGET)
document = Document(TARGET)

# Insert the new section immediately before the existing 기타사항 heading.
body = document._element.body
other_heading = next(
    paragraph._p
    for paragraph in document.paragraphs
    if paragraph.text.strip() == '5. 기타사항'
)
other_heading_text = other_heading.xpath('.//w:t')[0]
other_heading_text.text = '6. 기타사항'

new_heading = deepcopy(other_heading)
new_heading.xpath('.//w:t')[0].text = '5. 기타활동'

activity_table = deepcopy(document.tables[3]._tbl)
activity_table.xpath('.//w:t')
rows = activity_table.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
cells = rows[0].findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
header_values = ['일자', '수상·활동명', '주관기관', '비고']
for cell, value in zip(cells, header_values):
    cell_texts = cell.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
    if cell_texts:
        cell_texts[0].text = value
        for extra in cell_texts[1:]:
            extra.text = ''

value_cells = rows[1].findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
values = [
    '2026.05 ~ 2026.07',
    '기업 홈페이지 개선 및 연구데이터 분석 프로젝트 수행(팀장)',
    '(주)이에이포스',
    '고용노동부 청년일경험지원사업 프로젝트',
]
for cell, value in zip(value_cells, values):
    cell_texts = cell.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
    if cell_texts:
        cell_texts[0].text = value
        for extra in cell_texts[1:]:
            extra.text = ''

other_heading.addprevious(activity_table)
other_heading.addprevious(new_heading)
document.save(TARGET)
print(TARGET)