"""Offline geometry/font preview for editable native PowerPoint shapes."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import math,sys
from io import BytesIO

W=Path(__file__).resolve().parent
deck_path=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else W/'EOLWatch_요약서_김연동_기업관점.pptx'
W=deck_path.parent
deck=Presentation(deck_path)
scale=1800/deck.slide_width
points=scale*12700

def rgb(fill,default=None):
    try: return '#'+str(fill.fore_color.rgb)
    except (AttributeError,TypeError,ValueError): return default

def font(name,size,bold):
    if name=='Aptos':
        path='/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Aptos'+('-Bold' if bold else '')+'.ttf'
        return ImageFont.truetype(path,max(1,round(size*points)))
    return ImageFont.truetype('/System/Library/Fonts/AppleSDGothicNeo.ttc',max(1,round(size*points)),index=6 if bold else 0)

images=[]
overflow=[]
for index,slide in enumerate(deck.slides,1):
    im=Image.new('RGB',(1800,round(deck.slide_height*scale)),rgb(slide.background.fill,'#FAF9F9'))
    draw=ImageDraw.Draw(im)
    for sh in slide.shapes:
        x,y,w,h=[v*scale for v in [sh.left,sh.top,sh.width,sh.height]]
        box=(x,y,x+w,y+h)
        fill=rgb(sh.fill) if hasattr(sh,'fill') else None
        stroke=None
        try: stroke='#'+str(sh.line.color.rgb)
        except (AttributeError,TypeError,ValueError): pass
        lw=max(1,round((sh.line.width or 12700)*scale)) if hasattr(sh,'line') else 1
        kind=sh._element.tag.rsplit('}',1)[-1]
        if kind=='cxnSp':
            x1,y1,x2,y2=x,y,x+w,y+h
            transform=sh._element.spPr.xfrm
            if transform is not None and transform.get('flipH')=='1': x1,x2=x2,x1
            if transform is not None and transform.get('flipV')=='1': y1,y2=y2,y1
            draw.line((x1,y1,x2,y2),fill=stroke or '#C8C8C8',width=lw)
            if sh._element.xpath('.//a:tailEnd[@type="triangle"]'):
                angle=math.atan2(y2-y1,x2-x1);d=7
                draw.polygon([(x2,y2),(x2-d*math.cos(angle)+d*.42*math.sin(angle),y2-d*math.sin(angle)-d*.42*math.cos(angle)),(x2-d*math.cos(angle)-d*.42*math.sin(angle),y2-d*math.sin(angle)+d*.42*math.cos(angle))],fill=stroke)
        elif sh.shape_type==13:
            picture=Image.open(BytesIO(sh.image.blob)).convert('RGB')
            iw,ih=picture.size
            picture=picture.crop((round(iw*sh.crop_left),round(ih*sh.crop_top),round(iw*(1-sh.crop_right)),round(ih*(1-sh.crop_bottom))))
            picture=picture.resize((round(w),round(h)),Image.Resampling.LANCZOS)
            im.paste(picture,(round(x),round(y)))
        elif sh.shape_type==1:
            if sh.auto_shape_type==MSO_AUTO_SHAPE_TYPE.OVAL:
                draw.ellipse(box,fill=fill,outline=stroke,width=lw)
            elif sh.auto_shape_type==MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
                radius=min(w,h)*sh.adjustments[0]
                draw.rounded_rectangle(box,radius=radius,fill=fill,outline=stroke,width=lw)
            else:
                draw.rectangle(box,fill=fill,outline=stroke,width=lw)
        if not sh.has_text_frame or not sh.text: continue
        frame=sh.text_frame
        lines=[]
        for par in frame.paragraphs:
            if not par.runs: continue
            f=par.runs[0].font
            ft=font(f.name,f.size.pt,bool(f.bold))
            try: fc='#'+str(f.color.rgb)
            except: fc='#333333'
            spacing=par.line_spacing if isinstance(par.line_spacing,float) else 1.15
            height=f.size.pt*points*spacing
            text=par.text
            if draw.textlength(text,font=ft)>w+2: overflow.append((index,text,'width'))
            lines.append((text,ft,fc,height,par.alignment))
        total=sum(v[3] for v in lines)
        if total>h+2: overflow.append((index,sh.text,'height'))
        yy=y
        if frame.vertical_anchor==MSO_ANCHOR.MIDDLE: yy+=max(0,(h-total)/2)
        for text,ft,fc,lh,alignment in lines:
            tw=draw.textlength(text,font=ft)
            xx=x+(w-tw)/2 if alignment==PP_ALIGN.CENTER else x+w-tw if alignment==PP_ALIGN.RIGHT else x
            draw.text((xx,yy),text,font=ft,fill=fc,anchor='lt')
            yy+=lh
    im.save(W/f'layout-{index}.png')
    images.append(im)
sheet=Image.new('RGB',(1840,1070),'#C8C8C8')
for i,im in enumerate(images):
    im=im.copy();im.thumbnail((890,501))
    sheet.paste(im,(20+(i%2)*910,20+(i//2)*530))
sheet.save(W/'layout-review.png')
print('Potential text box issues:',overflow)
print(W/'layout-review.png')
