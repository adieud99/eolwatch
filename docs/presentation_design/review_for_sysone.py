"""Revise the brief after a company-side content review: cover + 3 body slides."""
from pathlib import Path
from copy import deepcopy
from shutil import copy2
from io import BytesIO
import hashlib, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent.parent
WORK=HERE/'sysone_review'
WORK.mkdir(exist_ok=True)
SOURCE=Path(json.loads((HERE/'corporate'/'delivered_files.json').read_text())[0])
source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
OUT=WORK/'EOLWatch_요약서_김연동_시스원활용제안.pptx'
copy2(SOURCE,OUT)

helper={'__file__':str(HERE/'design_presentation.py')}
exec((HERE/'design_presentation.py').read_text().split('for s in prs.slides:')[0],helper)
for name in ['txt','rect','ellipse','line','color','BG','INK','GREEN','GRAY','LINE','WHITE','KOREAN','LATIN']:
    globals()[name]=helper[name]
prs=Presentation(OUT)
for slide in prs.slides:
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    slide._element.set('showMasterSp','0')
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=color(BG)

def footer(sl,no):
    line(sl,.65,6.99,12.68,6.99,LINE,.6)
    txt(sl,.65,7.12,9,.18,'EOLWatch  /  시스원 업무 활용 제안',8.8,GRAY)
    txt(sl,11.8,7.1,.88,.23,f'{no:02d} / 04',10,GRAY,font=LATIN,align=PP_ALIGN.RIGHT)

def header(sl,section,title,subtitle):
    txt(sl,.65,.47,11.7,.23,section,10.5,GREEN,True)
    txt(sl,.65,1.0,12.05,.64,title,31.5,INK,True)
    txt(sl,.65,1.84,12.03,.43,subtitle,14.7,INK)

def cropped_picture(sl,path,x,y,w,crop):
    from PIL import Image
    iw,ih=Image.open(path).size
    l,t,r,b=crop
    h=w*(b-t)/(r-l)
    pic=sl.shapes.add_picture(str(path), Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    pic.crop_left=l/iw;pic.crop_top=t/ih;pic.crop_right=1-r/iw;pic.crop_bottom=1-b/ih
    pic.name='실제 웹 화면 — 분석 #3 → #4의 전후 비교'
    rect(sl,x,y,w,h,None,LINE,lw=.8)
    return pic

# Cover remains deliberately simple and uses the exact original palette.
s=prs.slides[0]
txt(s,.9,.72,11.4,.3,'시스원 유지관리·공개SW 기술지원 업무 활용 제안',12,GREEN,True)
txt(s,.84,2.35,11.55,1.26,'EOLWatch',68,INK,True,LATIN)
txt(s,.91,3.83,11.45,1.1,'고객사 자산의 취약점·지원 종료를\n담당자의 조치와 보고 근거로 연결합니다.',25,GREEN,spacing=1.25)
line(s,.91,6.02,12.43,6.02,LINE,.7)
txt(s,.91,6.35,9,.29,'한국폴리텍대학 광명융합기술교육원',13,INK)
txt(s,.91,6.82,9,.31,'데이터분석과  김연동',14,INK,True)
txt(s,10.5,6.83,1.93,.27,'2026.09.17',12,INK,font=LATIN,align=PP_ALIGN.RIGHT)
s.notes_slide.notes_text_frame.text='''EOLWatch는 등록한 고객사 서버와 프로젝트의 소프트웨어 구성, 취약점, 지원 종료 일정을 담당자의 조치 업무와 보고 자료로 연결하는 웹 서비스입니다.
시스원의 공개 사업 소개에 있는 시스템통합 유지관리 및 공개SW 기술지원 업무를 활용 대상으로 잡았습니다. 회사 내부에서 현재 어떤 시스템을 쓰는지, 어떤 문제가 발생하는지는 확인한 사실로 전제하지 않습니다.
이 자료는 개인 경험과 기술 나열 대신, 담당자가 수행할 업무와 제공받을 자료, 실제 구현 근거와 적용 범위를 설명합니다. 표지와 본문 3장으로 구성했습니다.
기업 사업 범위 출처: https://www.sysone.co.kr/businesses_veiw.php?idx=6 및 https://www.sysone.co.kr/businesses.php (2026-09-17 확인). 시스원 공식 제품이나 기존 업무 시스템과 연동을 완료했다는 의미는 아닙니다.'''

# Body 1: company use cases, decisions, and the exact information supplied.
s=prs.slides[1]
header(s,'01 / 업무 활용','유지관리 업무에서 무엇을 판단할 수 있는가',
       '활용 대상: 시스원의 시스템통합 유지관리·공개SW 기술지원 담당자')
cols=[.86,3.60,8.51]
widths=[2.47,4.51,3.96]
labels=['업무 상황','EOLWatch가 제공할 자료','담당자의 활용']
rect(s,.65,2.60,12.03,.56,GREEN,radius=True)
for x,w,label in zip(cols,widths,labels):
    txt(s,x,2.74,w,.29,label,13,WHITE,True)
rows=[
 ('취약점 공지 대응','해당 구성요소·버전과 연결된\n고객사·사이트·서버 목록',
  '확인할 대상을 찾고\n담당자에게 조치 업무 배정'),
 ('지원 종료 대응','지원 종료일·공식 출처와\n연결 자산·계약 정보',
  '업그레이드·교체 대상과\n작업 일정 검토'),
 ('조치 후 고객 보고','담당자·기한·처리 이력과\n동일 범위 재분석 비교 PDF',
  '처리 결과를 인계하고\n고객에게 조치 근거 전달'),
]
for i,row in enumerate(rows):
    y=3.45+i*1.05
    txt(s,cols[0],y+.07,widths[0],.46,row[0],18.7,GREEN,True)
    txt(s,cols[1],y,widths[1],.75,row[1],17.2,INK,spacing=1.25)
    txt(s,cols[2],y,widths[2],.75,row[2],17.2,INK,spacing=1.25)
    line(s,.85,y+.85,12.47,y+.85,LINE,.7)
txt(s,.85,6.70,11.6,.22,'분석된 자산과 등록된 지원 종료 일정·근거를 기준으로 조회합니다.',10.5,GRAY)
footer(s,2)
s.notes_slide.notes_text_frame.text='''시스원이 공개한 ITO 사업의 시스템통합 유지관리, 공개SW 기술지원 업무에서 활용할 수 있는 세 가지 상황을 제시합니다. 시스원의 현재 업무 방식이나 내부 도구의 부족을 단정한 표가 아닙니다.
취약점 공지 대응: 등록하고 분석한 자산의 구성요소·버전·CVE를 조회해 어느 고객사와 서버를 검토해야 하는지 확인할 수 있습니다. 이 정보를 담당자와 기한을 정한 조치 작업목록으로 연결합니다. 아직 수집하지 않은 자산까지 자동으로 파악한다는 의미는 아닙니다.
지원 종료 대응: 제품의 종료일과 공식 출처, 연결 자산과 계약을 참고해 변경 대상을 검토합니다. 현재는 등록된 날짜와 근거를 바탕으로 설명합니다. 공개 EOL 개별 조회·미리보기·적용의 전체 브라우저 검증은 별도 항목으로 남아 있습니다.
조치 후 고객 보고: 담당자의 작업 메모와 변경 이력, 동일 자산·동일 범위의 재분석 비교 자료를 보존합니다. PDF를 고객 설명 자료로 활용하고, 담당자 변경 시 이전 처리 근거를 확인할 수 있습니다.
업무 시간·비용 절감 수치는 측정하지 않았습니다. 이 장의 가치는 제공 자료가 어떤 판단과 업무를 지원하는지에 있습니다.
사업 출처: https://www.sysone.co.kr/businesses_veiw.php?idx=6
구현 근거: docs/IMPLEMENTATION_STATUS.md, docs/PROFESSOR_PROJECT_GUIDE.md, 조치 작업목록·제품·계약·비교 보고서 코드.'''

# Body 2: one real, bounded case and its original web evidence.
s=prs.slides[2]
header(s,'02 / 실제 검증 사례','업데이트 후 달라진 결과를 근거로 남깁니다',
       '동일 서버·동일 Python 환경을 재분석한 실제 웹 화면입니다.')
txt(s,.65,2.64,3.80,.29,'LAB-VM-01 · 지정 데모 환경',11.7,GREEN,True)
txt(s,.65,3.19,3.8,.43,'Jinja2 업데이트',23,INK,True)
txt(s,.65,3.91,3.86,.57,'3.1.4 → 3.1.6',29,GREEN,True,LATIN)
line(s,.65,4.78,4.05,4.78,LINE,.8)
txt(s,.65,5.08,3.8,.39,'재분석 미검출 3건',20,INK,True)
txt(s,.65,5.72,3.84,.72,'조치 기록 1건 저장·재조회\n비교 PDF·JSON 다운로드 확인',13.9,INK,spacing=1.30)
txt(s,4.65,2.64,8.03,.29,'실제 웹 화면 ①  |  분석 #3 → #4의 버전 변화·재분석 결과',11.7,GREEN,True)
cropped_picture(s,ROOT/'reports'/'demo-analysis-comparison.png',4.65,3.08,8.03,(316,574,1534,866))
txt(s,4.65,5.18,8.03,.29,'실제 웹 화면 ②  |  담당자의 조치 기록과 재분석 근거',11.7,GREEN,True)
cropped_picture(s,ROOT/'reports'/'vulnerability-action-history.png',4.65,5.60,8.03,(318,360,1533,556))
txt(s,.65,6.66,4.18,.21,'검증일 2026.09.15 · 조치 판단은 담당자가 기록',9.2,GRAY)
footer(s,3)
s.notes_slide.notes_text_frame.text='''이 장은 회사 도입 성과가 아니라 현재 구현한 조치·재점검 흐름의 실제 로컬 검증 사례입니다. 웹 화면은 reports/demo-analysis-comparison.png의 비교 패널과 reports/vulnerability-action-history.png의 조치 기록을 PowerPoint의 자르기 기능으로 표시했습니다. 원본 이미지 내용과 수치는 수정하지 않았습니다.
범위: LAB-VM-01의 지정 Python 데모 가상환경, 분석 #3 → #4, SBOM #12 → #13. 2026-09-15에 담당자가 Jinja2 3.1.4를 3.1.6으로 업데이트하고 서비스 정상 응답을 확인한 뒤 동일 범위를 재분석했습니다.
분석 결과: Jinja2 관련 CVE 3건이 재분석 미검출로 분류됐으며, 계속 검출·신규 검출·구성요소 제거는 각각 0건이었습니다. 이는 지정한 데모 범위와 당시 도구·DB 기준의 결과입니다.
조치 기록: CVE-2025-27516 한 건에 담당자가 재분석 근거와 작업 메모를 첨부해 수동 조치 완료를 기록했고 저장·재조회를 확인했습니다. 나머지 두 건까지 자동으로 완료 처리된 것은 아닙니다.
비교 PDF와 JSON 다운로드도 확인했습니다. 회사 담당자는 이처럼 실제 작업과 재점검 결과를 함께 보존하고 고객 설명이나 인계에 활용할 수 있습니다.
구성요소 식별·취약점 대조는 Syft·Grype를 사용합니다. 제가 구현한 가치는 분석 원본을 자산에 연결하고, 담당자 이력과 동일 범위 비교·보고서로 이어 주는 관리 기능입니다.
근거: reports/eolwatch-analysis-3-4.json, reports/vulnerability-action-history.png, docs/archive/REMEDIATION_VERIFICATION_2026-09-15.md, docs/archive/COMPARISON_REPORT_VERIFICATION_2026-09-15.md.'''

# Body 3: what the company receives and a realistic scope for a pilot.
s=prs.slides[3]
header(s,'03 / 제공물과 적용 조건','제공 결과물과 적용 전 확인사항',
       '분석 결과를 담당자의 조치 업무에 연결하는 웹 서비스와 운영 자료를 제공합니다.')
txt(s,.65,2.62,6.30,.32,'제공하는 결과물',13,GREEN,True)
deliverables=[
 ('웹 관리 화면','고객사 자산·구성요소·취약점·지원 종료 조회',3.23),
 ('조치·비교 자료','담당자·기한·이력 + 전후 비교 PDF·JSON',4.20),
 ('실행·운영 자료','소스 코드·Docker 실행 구성·백업/복구 안내',5.17),
]
for title,detail,y in deliverables:
    txt(s,.65,y,6.55,.38,title,20,INK,True)
    txt(s,.65,y+.48,6.6,.30,detail,14.0,INK)
    line(s,.65,y+.86,7.13,y+.86,LINE,.65)
txt(s,.65,6.43,6.48,.25,'직접 구현: 자산 연결 · 조치 이력 · 재분석 비교 · 보고서',11.8,GREEN,True)
line(s,7.53,2.58,7.53,6.72,LINE,.7)
rect(s,7.90,2.61,4.78,1.82,GREEN,radius=True)
txt(s,8.23,2.91,4.13,.24,'현재 로컬 검증 범위',11.5,WHITE,True)
txt(s,8.23,3.39,4.13,.39,'관리 VM 1대 + 대상 서버 2대',19.0,WHITE,True)
txt(s,8.23,4.00,4.13,.28,'SSH·ZIP 분석, 조치 기록·비교 보고서',12.4,WHITE)
txt(s,7.9,4.92,4.76,.31,'시범 적용 시 함께 확인할 항목',14.0,GREEN,True)
txt(s,7.9,5.51,4.76,1.08,'수집할 서버·프로젝트의 범위\nSSH 접근 권한과 데이터 보관 기준\n담당자 역할과 고객 보고 항목',15.1,INK,spacing=1.43)
footer(s,4)
s.notes_slide.notes_text_frame.text='''기업 담당자가 검토할 수 있는 제공물과 시범 적용 조건을 구분해 마무리합니다.
제공물은 웹 관리 화면, 분석 원본과 담당자 조치 기록·비교 보고서, 소스 코드와 Docker 실행 구성 및 백업·복구 안내입니다. SBOM은 소프트웨어 구성요소 명세서이며 분석 경로는 SPDX 2.3 형식을 사용합니다.
직접 구현한 범위는 자산·분석 범위 연결, 작업과 원본 저장, 담당자·기한·이력 관리, 동일 범위 전후 비교, PDF·JSON 보고서입니다. 분석 도구는 Syft·Grype를 활용하고 OSV 조회는 선택적인 보강에 사용합니다. 웹/API/DB는 React·FastAPI·PostgreSQL로 구성했습니다.
현재 확인한 환경은 관리 VM 1대와 대상 서버 2대의 로컬 환경입니다. ZIP 및 SSH 분석, 예약 분석, 실패 후 재시도, 조치 기록, 비교 보고서, DB·참조 ZIP의 수동 백업·임시 복원 검증 근거가 있습니다. 모든 확장 기능의 실제 브라우저 검증이나 기업 운영 환경의 성능·권한 적합성 검증을 완료했다는 의미는 아닙니다.
시범 적용을 검토한다면 수집을 허용할 서버와 프로젝트 범위, SSH 권한, 원본 및 조치 자료의 보관 기준, 담당자 역할과 고객 보고 양식을 먼저 맞춥니다. 실제 기업 적용 성과나 업무 시간 절감 수치는 아직 측정하지 않았습니다.
적용 후에는 대응 자산 식별의 정확성, 고객 보고에 필요한 항목의 충족 여부, 점검·자료 작성에 걸리는 시간을 확인하는 방식으로 활용성을 평가할 수 있습니다. 이는 제안하는 확인 기준이며 이미 달성한 성과 수치는 아닙니다.
기존 시스원 시스템과의 자동 연동, 고객사별 외부 로그인 격리, 자동 패치는 현재 제공 범위로 주장하지 않습니다. 이 장에서 기업의 검토 기준을 보여주되 이미 도입을 확정한 것처럼 표현하지 않습니다.'''

# Remove template effects and ensure Korean text has an explicit Korean typeface.
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    if any('가'<=c<='힣' for c in run.text): run.font.name=KOREAN
        sppr=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if sppr is not None:
            for child in list(sppr):
                if child.tag.rsplit('}',1)[-1] in ['effectLst','effectDag']: sppr.remove(child)
            sppr.append(OxmlElement('a:effectLst'))
        for el in sh._element.xpath('./p:style/a:effectRef'): el.set('idx','0')
prs.core_properties.title='EOLWatch — 시스원 업무 활용 제안'
prs.core_properties.subject='유지관리 업무 활용, 실제 조치 사례, 제공 결과물과 적용 범위'
prs.core_properties.author='김연동'
prs.save(OUT)

# Validate content, package structure and unchanged source.
check=Presentation(OUT)
assert len(check.slides)==4
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
allowed={'FAF9F9','4F785F','333333','A39D9D','C8C8C8'}
content=[]
for i,s in enumerate(check.slides,1):
    assert set(s._element.xpath('.//a:srgbClr/@val'))<=allowed
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=check.slide_width+100 and sh.top+sh.height<=check.slide_height+100,(i,sh.name)
    content.append({'slide':i,'text':[sh.text for sh in s.shapes if sh.has_text_frame and sh.text],
                    'notes':s.notes_slide.notes_text_frame.text})
(WORK/'content_review.json').write_text(json.dumps(content,ensure_ascii=False,indent=2))
(WORK/'validation.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':source_hash,
    'output':str(OUT),'slides':4,'body_slides':3,'source_unchanged':True,
    'evidence_image':'reports/demo-analysis-comparison.png'},ensure_ascii=False,indent=2))
print(OUT)
print('Saved cover + 3 body slides, source preserved, actual evidence image embedded.')
