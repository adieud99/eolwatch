"""Reuse original template slides 3/4 as pages 2/4, preserving approved pages 1/3."""
from copy import deepcopy
from pathlib import Path, PurePosixPath
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import posixpath
import unicodedata
from lxml import etree as E
from pptx import Presentation

HERE = Path(__file__).resolve().parent
WORK = HERE / 'base_revision'
WORK.mkdir(exist_ok=True)
SOURCE = Path(json.loads((HERE / 'sysone_review/delivered_files.json').read_text())[0])
BASE = next(p for p in SOURCE.parent.glob('*.pptx')
            if unicodedata.normalize('NFC', p.name) == 'ppt베이스4.pptx')
OUT = WORK / 'EOLWatch_요약서_김연동_베이스반영.pptx'
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
RNS = 'http://schemas.openxmlformats.org/package/2006/relationships'
CNS = 'http://schemas.openxmlformats.org/package/2006/content-types'

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

hashes = {str(p): sha(p) for p in (SOURCE, BASE)}
with ZipFile(SOURCE) as z:
    original = {n: z.read(n) for n in z.namelist()}
with ZipFile(BASE) as z:
    template = {n: z.read(n) for n in z.namelist()}
parts = dict(original)

def xml(b):
    return E.fromstring(b)

def encode(el):
    return E.tostring(el, xml_declaration=True, encoding='UTF-8', standalone=True)

current_size = xml(parts['ppt/presentation.xml']).find('p:sldSz', NS)
base_size = xml(template['ppt/presentation.xml']).find('p:sldSz', NS)
ratio = int(current_size.get('cx')) / int(base_size.get('cx'))
assert abs(ratio - int(current_size.get('cy')) / int(base_size.get('cy'))) < 1e-7

def scale_slide(root):
    # Scaling group extents AND their child coordinate system keeps group ratios.
    # Custom path coordinates remain intact within their own coordinate system.
    for xfrm in root.findall('.//a:xfrm', NS):
        for el in xfrm:
            if E.QName(el).localname in ('off', 'ext', 'chOff', 'chExt'):
                for key in ('x', 'y', 'cx', 'cy'):
                    if el.get(key) is not None:
                        el.set(key, str(round(int(el.get(key)) * ratio)))
    scale_attrs = {
        'rPr': ('sz', 'spc', 'kern'), 'defRPr': ('sz', 'spc', 'kern'),
        'endParaRPr': ('sz', 'spc', 'kern'), 'spcPts': ('val',),
        'pPr': ('marL', 'marR', 'indent'),
        'bodyPr': ('lIns', 'rIns', 'tIns', 'bIns', 'spcCol'),
        'ln': ('w',), 'tab': ('pos',),
        'outerShdw': ('blurRad', 'dist'), 'innerShdw': ('blurRad', 'dist'),
    }
    for el in root.iter():
        for key in scale_attrs.get(E.QName(el).localname, ()):
            if el.get(key) is not None:
                el.set(key, str(round(int(el.get(key)) * ratio)))

def replace_text(root, shape_id, text):
    shape = root.xpath(f'.//p:sp[p:nvSpPr/p:cNvPr[@id="{shape_id}"]]', namespaces=NS)[0]
    body = shape.find('p:txBody', NS)
    old_p = body.find('a:p', NS)
    ppr = old_p.find('a:pPr', NS)
    first = old_p.find('a:r/a:rPr', NS)
    assert first is not None
    for p in list(body.findall('a:p', NS)):
        body.remove(p)
    # Keep the original style, font family, line height, color and text box.
    p = E.SubElement(body, f'{{{NS["a"]}}}p')
    if ppr is not None:
        p.append(deepcopy(ppr))
    for n, line in enumerate(text.split('\n')):
        if n:
            br = E.SubElement(p, f'{{{NS["a"]}}}br')
            br.append(deepcopy(first))
        run = E.SubElement(p, f'{{{NS["a"]}}}r')
        run.append(deepcopy(first))
        E.SubElement(run, f'{{{NS["a"]}}}t').text = line

TEXT = {
    2: {
        4: '시스원 업무 활용',
        5: '유지관리·공개SW 기술지원 담당자를 위한 활용 제안',
        21: '프로젝트의 역할',
        14: '고객사 자산의 구성요소·취약점·지원 종료 정보를 연결해, 담당자가 대응 대상을 확인하도록 지원합니다.\n조치 이력과 동일 범위의 재분석 결과는 담당자 인계와 고객 보고에 필요한 근거로 제공합니다.',
        22: '업무별 활용',
        18: '취약점 대응',
        19: '지원 종료 대응',
        20: '조치 결과 보고',
        15: '구성요소·버전과 연결된\n고객사·서버를 확인하고\n담당자·기한을 지정합니다.',
        16: '공식 종료일·연결 자산과\n계약 정보를 함께 살펴\n업그레이드·교체를 검토합니다.',
        17: '담당자·처리 이력과 재분석\n비교 PDF를 모아 조치 근거를\n고객 보고와 인계에 활용합니다.',
    },
    4: {
        10: '제공물과 적용 조건',
        11: '시스원 담당자가 활용할 결과물과 적용 전 확인 사항',
        20: '제공 결과물',
        21: '적용 전 확인',
        12: '웹 관리 화면: 고객사 자산·구성요소·취약점·지원 종료 조회\n조치·보고 자료: 담당자·기한·처리 이력, 전후 비교 PDF·JSON\n실행·운영 자료: 소스 코드·Docker 구성·백업·복구 안내',
        13: '점검 범위: 대상 서버·프로젝트와 분석 주기\n접근·보관 기준: SSH 권한과 분석 원본·조치 기록 보관\n업무 기준: 담당자 역할·조치 완료 판단·고객 보고 항목',
    },
}

def import_slide(target_no, base_no):
    root = xml(template[f'ppt/slides/slide{base_no}.xml'])
    target_rels_name = f'ppt/slides/_rels/slide{target_no}.xml.rels'
    rels = xml(original[target_rels_name])
    # Preserve the approved deck's layout and presenter notes only.
    for rel in list(rels):
        if not rel.get('Type').endswith(('/slideLayout', '/notesSlide')):
            rels.remove(rel)
    source_rels = xml(template[f'ppt/slides/_rels/slide{base_no}.xml.rels'])
    used = {value for el in root.iter() for key, value in el.attrib.items()
            if key.startswith('{' + NS['r'] + '}')}
    mapping = {}
    for rel in source_rels:
        old_id = rel.get('Id')
        if old_id not in used:
            continue
        assert rel.get('Type').endswith('/image'), rel.attrib
        src_path = posixpath.normpath(posixpath.join('ppt/slides', rel.get('Target')))
        dest_path = 'ppt/media/base_reuse_' + PurePosixPath(src_path).name
        parts[dest_path] = template[src_path]
        new_id = f'rIdBase{base_no}_{old_id}'
        mapping[old_id] = new_id
        nr = E.SubElement(rels, f'{{{RNS}}}Relationship')
        nr.attrib.update({'Id': new_id, 'Type': rel.get('Type'),
                          'Target': posixpath.relpath(dest_path, 'ppt/slides')})
    for el in root.iter():
        for key, val in list(el.attrib.items()):
            if key.startswith('{' + NS['r'] + '}'):
                el.set(key, mapping[val])
    root.set('showMasterSp', '0')
    for shape_id, text in TEXT[target_no].items():
        replace_text(root, shape_id, text)
    scale_slide(root)
    parts[f'ppt/slides/slide{target_no}.xml'] = encode(root)
    parts[target_rels_name] = encode(rels)
    # Append provenance to the existing factual notes without changing their scope.
    note_rel = next(r for r in rels if r.get('Type').endswith('/notesSlide'))
    note_path = posixpath.normpath(posixpath.join('ppt/slides', note_rel.get('Target')))
    note_root = xml(parts[note_path])
    body = note_root.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:txBody', namespaces=NS)[0]
    p = E.SubElement(body, f'{{{NS["a"]}}}p')
    r = E.SubElement(p, f'{{{NS["a"]}}}r')
    E.SubElement(r, f'{{{NS["a"]}}}t').text = (
        f'2026-09-17 수정: 사용자가 제공한 ppt베이스4.pptx의 {base_no}번 슬라이드를 직접 복사했습니다. '
        '도형·아이콘·색상·원본 서체·배치를 유지하고 화면 크기에 맞춰 전체를 등비 축소했습니다. '
        '본문 문구만 교체했으며 표지와 승인된 3페이지는 그대로 유지했습니다.'
    )
    parts[note_path] = encode(note_root)

import_slide(2, 3)
import_slide(4, 4)

# Register any SVG/raster content types that the approved deck did not contain.
ct = xml(parts['[Content_Types].xml'])
known = {el.get('Extension') for el in ct if E.QName(el).localname == 'Default'}
for item in xml(template['[Content_Types].xml']):
    if E.QName(item).localname == 'Default' and item.get('Extension') not in known:
        ct.append(deepcopy(item))
parts['[Content_Types].xml'] = encode(ct)
with ZipFile(OUT, 'w', ZIP_DEFLATED) as z:
    for name, data in parts.items():
        z.writestr(name, data)

check = Presentation(OUT)
assert len(check.slides) == 4
with ZipFile(OUT) as z:
    assert z.testzip() is None
    for no in (1, 3):
        for name in (f'ppt/slides/slide{no}.xml', f'ppt/slides/_rels/slide{no}.xml.rels'):
            assert z.read(name) == original[name]
for p in (SOURCE, BASE):
    assert sha(p) == hashes[str(p)]
visible = []
for i, slide in enumerate(check.slides, 1):
    visible.append({'slide': i, 'text': [sh.text for sh in slide.shapes if sh.has_text_frame and sh.text]})
(WORK / 'content_review.json').write_text(json.dumps(visible, ensure_ascii=False, indent=2))
(WORK / 'validation.json').write_text(json.dumps({
    'source': str(SOURCE), 'base': str(BASE), 'output': str(OUT),
    'original_hashes': hashes, 'sources_unchanged': True,
    'preserved_slide_xml_and_relationships': [1, 3],
    'copied_base_slides': {'2': 3, '4': 4},
    'scale_ratio': ratio, 'template_fonts_colors_geometry_preserved': True,
}, ensure_ascii=False, indent=2))
print(OUT)
print('Validated: original sources unchanged; slides 1/3 byte-for-byte preserved; 2/4 copied from base.')
