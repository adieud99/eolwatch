"""Editable four-slide redesign; preserves the source deck's speaker notes."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
SOURCE = Path((HERE / 'source_path.txt').read_text())
OUTPUT = HERE / 'EOLWatch_요약서_김연동_디자인.pptx'
SOURCE_HASH = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
prs = Presentation(SOURCE)
original_notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

# Exact visible palette extracted from the user's ppt베이스4.pptx.
BG = 'FAF9F9'
INK = '333333'
DARK = '4F785F'
GREEN = '4F785F'
MUTED = '333333'
GRAY = 'A39D9D'
LINE = 'C8C8C8'
PALE = 'FAF9F9'
LIME = 'FAF9F9'
WHITE = 'FAF9F9'
SOFT = 'FAF9F9'
KOREAN = 'Apple SD Gothic Neo'
LATIN = 'Aptos'

def color(value):
    return RGBColor.from_string(value)

def rect(sl, x, y, w, h, fill=None, stroke=None, radius=False, lw=1):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        sh.adjustments[0] = min(.12, .12 / min(w, h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color(fill)
    else:
        sh.fill.background()
    if stroke:
        sh.line.color.rgb = color(stroke)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh

def ellipse(sl, x, y, w, h=None, fill=None, stroke=None, lw=1):
    sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h or w))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color(fill)
    else:
        sh.fill.background()
    if stroke:
        sh.line.color.rgb = color(stroke)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh

def line(sl, x1, y1, x2, y2, c=LINE, lw=1, arrow=False, dash=False):
    sh = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color(c)
    sh.line.width = Pt(lw)
    if arrow:
        el = OxmlElement('a:tailEnd')
        el.set('type', 'triangle')
        el.set('w', 'sm')
        el.set('len', 'sm')
        sh.line._get_or_add_ln().append(el)
    if dash:
        el = OxmlElement('a:prstDash')
        el.set('val', 'dash')
        sh.line._get_or_add_ln().append(el)
    return sh

def txt(sl, x, y, w, h, content, size=16, c=INK, bold=False, font=KOREAN,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spacing=1.15):
    sh = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = content.replace('\n', ' / ')[:90]
    tf = sh.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    for i, value in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = value
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color(c)
        rpr = r._r.get_or_add_rPr()
        rpr.set('lang', 'ko-KR')
        for tag in ['a:ea', 'a:cs']:
            n = OxmlElement(tag)
            n.set('typeface', KOREAN)
            rpr.append(n)
    return sh

def pill(sl, x, y, w, text, fill=PALE, c=GREEN, size=10):
    rect(sl, x, y, w, .31, fill, radius=True)
    txt(sl, x, y, w, .31, text, size, c, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

def footer(sl, no, dark=False):
    c = SOFT if dark else GRAY
    line(sl, .65, 6.96, 12.68, 6.96, LINE, .6)
    txt(sl, .65, 7.1, 5, .2, 'EOLWatch  /  프로젝트 요약서', 9, c)
    txt(sl, 11.6, 7.08, 1.08, .24, f'{no:02d} / 04', 10, c, font=LATIN, align=PP_ALIGN.RIGHT)

def header(sl, section, title, subtitle, dark=False):
    txt(sl, .65, .47, 7, .22, section, 10, LIME if dark else GREEN, True, LATIN)
    txt(sl, .65, .99, 11.8, .65, title, 34, WHITE if dark else INK, True)
    txt(sl, .65, 1.79, 11.8, .38, subtitle, 15, SOFT if dark else MUTED)

def icon(sl, kind, x, y, size=.52, c=GREEN):
    # All icons are editable native PowerPoint line/shape elements.
    def L(a,b,d,e): return line(sl,x+a*size,y+b*size,x+d*size,y+e*size,c,1.6)
    if kind == 'document':
        rect(sl,x+.18*size,y+.08*size,.64*size,.84*size,None,c,True,1.6)
        L(.32,.34,.68,.34); L(.32,.51,.68,.51); L(.32,.68,.57,.68)
    elif kind == 'search':
        ellipse(sl,x+.1*size,y+.08*size,.59*size,.59*size,None,c,1.6)
        L(.62,.62,.9,.92); L(.39,.24,.39,.43)
        ellipse(sl,x+.365*size,y+.49*size,.05*size,fill=c)
    elif kind == 'asset':
        rect(sl,x+.13*size,y+.04*size,.74*size,.28*size,None,c,True,1.6)
        rect(sl,x+.13*size,y+.4*size,.74*size,.28*size,None,c,True,1.6)
        L(.27,.18,.29,.18); L(.27,.54,.29,.54)
        L(.5,.7,.5,.91); L(.23,.91,.77,.91)
    elif kind == 'check':
        ellipse(sl,x+.08*size,y+.08*size,.84*size,.84*size,None,c,1.6)
        L(.27,.51,.44,.67); L(.44,.67,.75,.34)

for s in prs.slides:
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    s._element.set('showMasterSp', '0')
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(BG)

# 01 — cover: minimal typography with a single-line descriptor.
s = prs.slides[0]
txt(s,.9,.72,7,.3,'프로젝트 요약서',12,GREEN,True)
txt(s,.84,2.6,11.55,1.26,'EOLWatch',68,INK,True,LATIN)
txt(s,.91,4.0,11.45,.55,'SBOM · 취약점 · 조치 통합 관리',24,GREEN)
line(s,.91,6.02,12.43,6.02,LINE,.7)
txt(s,.91,6.35,9,.29,'한국폴리텍대학 광명융합기술교육원',13,MUTED)
txt(s,.91,6.82,9,.31,'데이터분석과  김연동',14,INK,True)
txt(s,10.5,6.83,1.93,.27,'2026.09.17',12,MUTED,font=LATIN,align=PP_ALIGN.RIGHT)

# 02 — origin: three experiences and one clear motivation.
s=prs.slides[1]
header(s,'01 / BACKGROUND','프로젝트 배경','운영 경험에서 출발한 취약점 관리의 필요성')
line(s,1,2.95,1,5.83,LINE,1.1)
items=[('01','NginX 방어','외부 공격 상황에서도 서비스를\n계속 살려 두는 AWS 팀 과제',2.64),
       ('02','StockCast 운영','재고관리 시스템을 직접 배포하고\n운영한 개인 프로젝트',3.99),
       ('03','Log4j 사례','취약점 공개 당시 현장에서\n무엇이 문제였는지 찾아본 사례',5.34)]
for num,title,body,y in items:
    ellipse(s,.7,y,.62,fill=BG,stroke=LINE)
    txt(s,.7,y,.62,.62,num,14,GREEN,True,LATIN,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    txt(s,1.57,y-.01,4.37,.39,title,21,INK,True)
    txt(s,1.57,y+.51,4.47,.65,body,15.8,MUTED,spacing=1.22)
rect(s,6.55,2.61,6.13,3.95,DARK,radius=True)
pill(s,6.91,2.95,1.75,'프로젝트를 시작한 계기',fill=BG,c=GREEN,size=10)
txt(s,6.93,3.55,5.37,1.45,'서버를 지키는 경험,\n취약점에 대비하는 관리로.',27,WHITE,True,spacing=1.24)
line(s,6.94,5.13,7.43,5.13,LIME,2)
txt(s,6.94,5.43,5.33,.88,'두 과제는 서버를 멈추지 않게 지키는 일이었다.\nLog4j 사례를 보며, 내부 취약점을 미리 파악하고\n대비할 필요성을 느꼈다.',14.5,SOFT,spacing=1.2)
footer(s,2)

# 03 — features: one left-to-right operational flow.
s=prs.slides[2]
s.background.fill.fore_color.rgb=color(BG)
header(s,'02 / CAPABILITIES','주요 기능','구성요소를 식별하고, 영향 자산을 찾아 조치 이력까지 연결합니다.')
xs=[.65,3.76,6.86,9.96]
stages=[('01','SBOM 생성','구성요소 · 버전 식별','설치된 구성요소와 버전을\n식별해 SPDX 2.3 형식으로\n남깁니다.','SPDX 2.3','document'),
        ('02','취약점 분석','CVE · 수정 버전 확인','알려진 취약점 데이터와\n대조해 CVE와 수정 버전을\n찾습니다.','CVE / FIX VERSION','search'),
        ('03','자산 연결','영향받는 서버 확인','분석 결과를 고객사·사이트·\n자산에 연결해 영향받는\n서버를 표시합니다.','CUSTOMER / SITE / ASSET','asset'),
        ('04','조치 관리','이력 · 재점검 비교','담당자·기한·상태와 이력을\n남기고 재분석으로\n조치 전후를 비교합니다.','OWNER / STATUS / HISTORY','check')]
line(s,.88,2.84,12.21,2.84,LINE,1,arrow=True)
for x,(num,title,kicker,body,tag,ic) in zip(xs,stages):
    ellipse(s,x,2.61,.45,fill=GREEN)
    txt(s,x,2.61,.45,.45,num,12,WHITE,True,LATIN,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    icon(s,ic,x,3.38,.55,GREEN)
    txt(s,x,4.15,2.8,.42,title,23,INK,True)
    txt(s,x,4.75,2.8,.28,kicker,13,GREEN,True)
    txt(s,x,5.23,2.85,.97,body,15.3,INK,spacing=1.24)
    txt(s,x,6.45,2.8,.22,tag,8.8,GRAY,True,LATIN)
footer(s,3)

# 04 — technical architecture: editable boxes with meaningful connectors.
s=prs.slides[3]
header(s,'03 / ARCHITECTURE','기술 구성','로컬 VM에서 수집·분석하고, 표준 SBOM과 조치 데이터를 함께 관리합니다.')
txt(s,.65,2.46,6.8,.25,'SYSTEM FLOW',10,GREEN,True,LATIN)
txt(s,8.36,2.46,4.2,.25,'TECH STACK',10,GREEN,True,LATIN)
line(s,7.99,2.43,7.99,6.61,LINE,.7)

# Browser and application boundary.
rect(s,.67,3.37,1.29,1.16,WHITE,LINE,True)
rect(s,1.03,3.57,.54,.37,None,GREEN,True,1)
line(s,1.03,3.66,1.57,3.66,GREEN,.8)
txt(s,.73,4.08,1.17,.27,'사용자 브라우저',12,INK,True,align=PP_ALIGN.CENTER)
line(s,1.97,3.97,2.23,3.97,GREEN,1.2,True)
rect(s,2.26,2.99,5.29,2.76,PALE,LINE,radius=True,lw=.8)
txt(s,2.45,3.15,4.9,.25,'EOLWatch  /  Docker Compose',11,GREEN,True,LATIN)
rect(s,2.48,3.63,2.13,.71,WHITE,LINE,True)
txt(s,2.51,3.83,2.07,.32,'React 19 · Nginx',14,INK,True,LATIN,PP_ALIGN.CENTER)
rect(s,5.0,3.63,2.13,.71,WHITE,LINE,True)
txt(s,5.03,3.83,2.07,.32,'FastAPI',15,INK,True,LATIN,PP_ALIGN.CENTER)
line(s,4.64,3.98,4.96,3.98,GREEN,1.2,True)
rect(s,5.0,4.73,2.13,.72,DARK,radius=True)
txt(s,5.05,4.8,2.03,.28,'worker',15,WHITE,True,LATIN,PP_ALIGN.CENTER)
txt(s,5.03,5.15,2.07,.22,'SSH 수집 · SBOM · 분석',9.8,SOFT,align=PP_ALIGN.CENTER)
line(s,6.07,4.36,6.07,4.67,GREEN,1.2,True)
rect(s,2.48,4.73,2.13,.72,WHITE,LINE,True)
txt(s,2.54,4.79,2.01,.28,'PostgreSQL 16',14,INK,True,LATIN,PP_ALIGN.CENTER)
txt(s,2.54,5.16,2.01,.2,'SBOM 원본 · 총 17개 테이블',9.5,MUTED,align=PP_ALIGN.CENTER)
line(s,4.95,5.09,4.66,5.09,GREEN,1.2,True)
# Worker connects directly to the inspected servers, not through the DB.
line(s,6.07,5.47,6.07,5.94,GREEN,1.2)
line(s,3.3,5.94,6.43,5.94,GREEN,1.2)
line(s,3.3,5.94,3.3,6.13,GREEN,1.2,True)
line(s,6.43,5.94,6.43,6.13,GREEN,1.2,True)
rect(s,4.13,5.79,1.44,.27,BG)
txt(s,4.13,5.79,1.44,.23,'SSH 22 · 키 인증',9,GREEN,True,align=PP_ALIGN.CENTER)
for x,n in [(2.49,1),(5.47,2)]:
    rect(s,x,6.19,1.94,.46,WHITE,LINE,True)
    txt(s,x+.06,6.26,1.82,.3,f'대상 서버 {n} · 로컬 VM',10.4,INK,True,align=PP_ALIGN.CENTER)

# Stack list: current analysis and future integration are visibly distinguished.
stack=[(2.98,'웹 / API','React 19 · FastAPI','화면 · API 서버'),
       (4.10,'저장 / 표준','PostgreSQL 16 · SPDX 2.3','관계 데이터 · SBOM 원본'),
       (5.22,'분석','OSV API','CVE · 수정 버전 조회')]
for y,label,main,sub in stack:
    txt(s,8.36,y,3.98,.22,label,10.5,GREEN,True)
    txt(s,8.36,y+.34,4.32,.37,main,18.8,INK,True,LATIN)
    txt(s,8.36,y+.77,4.2,.24,sub,11.5,MUTED)
    if y<5: line(s,8.36,y+1.04,12.68,y+1.04,LINE,.65)
pill(s,8.36,6.43,.81,'연동 검토',size=9)
txt(s,9.28,6.46,3.3,.24,'Syft · Grype',12,MUTED,font=LATIN)
footer(s,4)

# Override decorative effects inherited from the source template.
for slide in prs.slides:
    for sh in slide.shapes:
        sppr=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if sppr is not None:
            for tag in ['a:effectLst','a:effectDag']:
                for child in list(sppr.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])):
                    sppr.remove(child)
            sppr.append(OxmlElement('a:effectLst'))
        for effect in sh._element.xpath('./p:style/a:effectRef'):
            effect.set('idx','0')

# Apply a consistent editing theme without changing the speaker notes.
from lxml import etree
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
for part in prs.part.package.iter_parts():
    if '/theme/theme' in str(part.partname) and part.content_type.endswith('theme+xml'):
        root=etree.fromstring(part.blob)
        for e in root.xpath('//a:fontScheme/a:majorFont/a:latin | //a:fontScheme/a:minorFont/a:latin',namespaces=ns):
            e.set('typeface',LATIN)
        for e in root.xpath('//a:fontScheme/a:majorFont/a:ea | //a:fontScheme/a:minorFont/a:ea',namespaces=ns):
            e.set('typeface',KOREAN)
        part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
prs.core_properties.title='EOLWatch 프로젝트 요약서'
prs.core_properties.subject='SBOM · 취약점 · 조치 통합 관리'
prs.core_properties.author='김연동'
prs.core_properties.keywords='EOLWatch, SBOM, CVE, 자산 관리, 조치 관리'
prs.save(OUTPUT)

# Structural validation of the actual saved file.
check=Presentation(OUTPUT)
assert len(check.slides)==4
assert [s.notes_slide.notes_text_frame.text for s in check.slides]==original_notes
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_HASH
outside=[]
for i,sl in enumerate(check.slides,1):
    for sh in sl.shapes:
        if sh.left < -100 or sh.top < -100 or sh.left+sh.width > check.slide_width+100 or sh.top+sh.height > check.slide_height+100:
            outside.append((i,sh.name))
assert not outside,outside
(HERE/'validation.json').write_text(json.dumps({'slides':4,'speaker_notes_preserved':True,
    'source_unchanged':True,'source_sha256':SOURCE_HASH,'output':str(OUTPUT),
    'editable_shapes':[len(s.shapes) for s in check.slides]},ensure_ascii=False,indent=2))
print(OUTPUT)
print('Validated: 4 slides, original notes preserved, no off-slide shapes, original untouched.')
