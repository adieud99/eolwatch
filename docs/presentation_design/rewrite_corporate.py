"""Create a separate, business-facing brief from the approved design copy."""
from pathlib import Path
from shutil import copy2
from copy import deepcopy
import json
import hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
WORK = HERE / 'corporate'
WORK.mkdir(exist_ok=True)
BASE = Path(next(p for p in json.loads((HERE/'delivered_files.json').read_text()) if p.endswith('.pptx')))
OUT = WORK/'EOLWatch_요약서_김연동_기업관점.pptx'
base_hash = hashlib.sha256(BASE.read_bytes()).hexdigest()
copy2(BASE, OUT)

# Reuse the approved palette and drawing helpers without running the old builder.
helpers = {'__file__':str(HERE/'design_presentation.py')}
exec((HERE/'design_presentation.py').read_text().split('for s in prs.slides:')[0], helpers)
for name in ['txt','rect','ellipse','line','pill','footer','header','icon','color',
             'BG','INK','GREEN','GRAY','LINE','WHITE','KOREAN','LATIN']:
    globals()[name] = helpers[name]

prs = Presentation(OUT)
for s in prs.slides:
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    s._element.set('showMasterSp','0')
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(BG)

# 1. Lead with the work the company can do with the service.
s=prs.slides[0]
txt(s,.9,.72,11.4,.3,'시스원 유지보수 업무를 위한 제안',12,GREEN,True)
txt(s,.84,2.34,11.55,1.26,'EOLWatch',68,INK,True,LATIN)
txt(s,.91,3.81,11.45,1.12,'고객사 자산의 취약점·지원 종료·조치를\n한곳에서 관리하는 웹 서비스',25,GREEN,spacing=1.24)
line(s,.91,6.02,12.43,6.02,LINE,.7)
txt(s,.91,6.35,9,.29,'한국폴리텍대학 광명융합기술교육원',13,INK)
txt(s,.91,6.82,9,.31,'데이터분석과  김연동',14,INK,True)
txt(s,10.5,6.83,1.93,.27,'2026.09.17',12,INK,font=LATIN,align=PP_ALIGN.RIGHT)
s.notes_slide.notes_text_frame.text = '''EOLWatch는 고객사 자산의 소프트웨어 구성과 취약점, 지원 종료 일정, 담당자 조치 이력을 한곳에서 관리하는 웹 서비스입니다.
시스원의 유지보수 업무에 활용할 수 있는 기능과 산출물을 중심으로 제안합니다. 취약점 발견 이후 어느 자산을 확인하고, 누가 언제 조치하며, 고객에게 어떤 근거를 전달할지를 연결하는 것이 목적입니다.
제가 제공하는 부분은 분석 결과와 고객사 자산을 연결하는 관리 화면, 담당자별 조치 업무, 재분석 비교와 보고서입니다. 구성요소 분석에는 Syft, 취약점 대조에는 Grype를 사용합니다.
이 자료는 현재 구현한 기능의 활용 제안입니다. 시스원 내부 시스템과 연계하거나 고객사 운영 환경에 실제 도입한 성과를 주장하는 자료는 아닙니다.'''

# 2. Company-facing scenarios, without inventing the company's current problems.
s=prs.slides[1]
header(s,'01 / 업무 과제','유지보수 담당자가 확인해야 할 세 가지',
       '영향 대상과 지원 종료 일정을 확인하고, 조치 결과를 고객에게 설명할 수 있어야 합니다.')
line(s,1,2.96,1,5.87,LINE,1.1)
items=[
 ('01','어느 고객사·서버가 영향을 받는가',
  '알려진 취약점이 있는 구성요소와 버전을\n고객사·사이트·서버에 연결해 확인합니다.',2.66),
 ('02','무엇을 미리 교체·업그레이드할 것인가',
  '지원 종료일과 영향 자산을 함께 보고\n변경 대상과 작업 시점을 검토합니다.',4.02),
 ('03','누가 처리했고, 결과는 어떻게 달라졌는가',
  '담당자·기한·처리 내용을 기록하고\n재분석 결과와 보고서로 확인합니다.',5.38),
]
for num,title,body,y in items:
    ellipse(s,.7,y,.62,fill=BG,stroke=LINE)
    txt(s,.7,y,.62,.62,num,14,GREEN,True,LATIN,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    txt(s,1.55,y+.01,4.7,.45,title,17.8,INK,True)
    txt(s,1.55,y+.55,4.7,.66,body,14.7,INK,spacing=1.25)
rect(s,6.68,2.65,6.0,3.95,GREEN,radius=True)
txt(s,7.06,3.03,5.22,.31,'이 프로젝트가 제공할 가치',12,WHITE,True)
txt(s,7.06,3.69,5.24,1.28,'대상 확인부터\n조치·고객 보고까지 연결',26,WHITE,True,spacing=1.24)
line(s,7.07,5.25,7.6,5.25,WHITE,1.7)
txt(s,7.07,5.54,5.17,.75,'고객사 자산·분석 결과·처리 이력을 묶어\n담당자 간 인계와 고객 보고에 필요한\n근거를 제공합니다.',14.3,WHITE,spacing=1.23)
footer(s,2)
s.notes_slide.notes_text_frame.text = '''기업 담당자가 EOLWatch로 수행할 수 있는 업무를 세 가지로 설명합니다.
첫째, 취약점 대응입니다. 등록하고 분석한 자산에서 구성요소 이름과 버전을 확인하고 고객사·사이트·서버와 연결해 대응 대상을 찾습니다. 산출물은 영향 자산과 구성요소 목록입니다.
둘째, 지원 종료 대응입니다. 제품의 지원 종료일과 공식 출처, 연결된 자산과 계약을 확인해 업그레이드 또는 교체 검토에 활용합니다. 날짜는 등록된 값과 근거를 기준으로 하며 모든 제품에 자동으로 최신 일정을 적용한다고 설명하지 않습니다.
셋째, 조치와 보고입니다. 담당자·기한·상태·실제 작업 메모를 남기고 같은 자산과 범위를 다시 분석해 전후 결과를 비교합니다. 작업목록과 비교 보고서를 인계 및 고객 설명 자료로 활용할 수 있습니다.
시스원이 현재 수작업을 하고 있다거나 특정 문제를 겪는다고 단정하지 않습니다. 시간·비용 절감 수치는 측정하지 않았으므로 정량 성과를 제시하지 않습니다.'''

# 3. State deliverables alongside each capability.
s=prs.slides[2]
header(s,'02 / 제공 기능','담당자에게 제공하는 기능과 결과물',
       '등록한 자산의 현황부터 조치 근거까지, 업무에 필요한 자료를 연결해 제공합니다.')
xs=[.65,3.76,6.86,9.96]
stages=[
 ('01','자산·구성 파악','무엇이 설치되어 있는지',
  '고객사·서버별 구성요소와\n버전을 수집하고\n자산에 연결합니다.',
  '자산별 구성 목록 · SBOM','document'),
 ('02','위험·기한 확인','무엇을 먼저 확인할지',
  '알려진 취약점·수정 버전과\n지원 종료일을 함께 보고\n대응 우선순위를 판단합니다.',
  'CVE · 지원 종료 대상 목록','search'),
 ('03','조치 업무 관리','누가 언제 처리할지',
  '담당자·기한·진행 상태와\n실제 처리 내용,\n변경 이력을 기록합니다.',
  '담당자별 조치 작업목록','asset'),
 ('04','재점검·보고','처리 결과가 어떻게 바뀌었는지',
  '같은 범위의 재분석으로\n전후 변화를 비교하고\n보고서로 내려받습니다.',
  '비교 보고서 · PDF / JSON','check'),
]
line(s,.88,2.84,12.21,2.84,LINE,1,arrow=True)
for x,(num,title,kicker,body,result,ic) in zip(xs,stages):
    ellipse(s,x,2.61,.45,fill=GREEN)
    txt(s,x,2.61,.45,.45,num,12,WHITE,True,LATIN,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    icon(s,ic,x,3.38,.55,GREEN)
    txt(s,x,4.13,2.86,.44,title,22,INK,True)
    txt(s,x,4.75,2.85,.33,kicker,12.3,GREEN,True)
    txt(s,x,5.23,2.85,1.0,body,14.8,INK,spacing=1.25)
    line(s,x,6.35,x+2.62,6.35,LINE,.65)
    txt(s,x,6.5,2.85,.25,result,10.5,GREEN,True)
footer(s,3)
s.notes_slide.notes_text_frame.text = '''기능 이름보다 담당자가 얻게 되는 자료를 중심으로 설명합니다.
자산·구성 파악: 등록 서버의 지정 범위를 SSH로 수집하거나 소스 ZIP·기존 SBOM을 반입합니다. 구성요소 목록을 고객사·사이트·서버 자산과 연결합니다. SBOM은 소프트웨어 구성요소와 버전을 담은 명세서입니다.
위험·기한 확인: 알려진 취약점(CVE)과 수정 버전, 등록된 지원 종료 일정을 확인합니다. EOL은 제품의 지원 종료를 뜻합니다. CVE 심각도, 지원 종료 임박도 등을 함께 보고 담당자가 대응 순서를 정할 수 있도록 합니다.
조치 업무 관리: 담당자, 기한, 진행 상태, 작업 메모와 재분석 근거를 관리합니다. 이전 분석에서 남은 미완료 조치도 작업목록으로 확인할 수 있습니다.
재점검·보고: 동일 자산과 분석 범위의 전후 결과를 비교하고, 계속 검출·신규 검출·재분석 미검출·구성요소 제거를 구분해 PDF와 JSON으로 제공합니다.
실제 업데이트와 서비스 정상 동작 확인, 조치 완료 판단은 담당자가 수행합니다. 탐지되지 않았다는 결과만으로 자동 완료 처리하지 않습니다.'''

# 4. Show a practical application path and bounded implementation evidence.
s=prs.slides[3]
header(s,'03 / 적용 방식','적용 방식과 구현 근거',
       '대상 범위를 정해 분석하고, 담당자가 조치한 결과를 재점검과 보고서로 남깁니다.')
txt(s,.65,2.46,7.0,.28,'업무 적용 흐름',12,GREEN,True)
boxes=[
 (.65,'대상 등록','고객사·서버 자산\n서버 경로 또는 소스 ZIP'),
 (3.1,'분석·관리','구성요소·취약점 확인\n지원 종료·담당자 연결'),
 (5.55,'조치·재점검','담당자가 업데이트\n동일 범위 재분석·보고'),
]
for x,title,body in boxes:
    rect(s,x,2.99,2.15,1.4,BG,LINE,True)
    txt(s,x+.13,3.21,1.89,.35,title,17.5,GREEN,True,align=PP_ALIGN.CENTER)
    txt(s,x+.1,3.79,1.95,.46,body,11.8,INK,align=PP_ALIGN.CENTER,spacing=1.23)
line(s,2.82,3.7,3.07,3.7,GREEN,1.1,True)
line(s,5.27,3.7,5.52,3.7,GREEN,1.1,True)
txt(s,.65,4.77,7.05,.27,'구현 구성',12,GREEN,True)
stack=[('분석','Syft → SPDX 2.3 → Grype',5.25),
       ('웹·저장','React · FastAPI · PostgreSQL',5.79),
       ('운영','Docker Compose · 예약 분석 · DB/ZIP 백업',6.33)]
for label,value,y in stack:
    txt(s,.65,y,1.1,.28,label,12.5,GREEN,True)
    txt(s,1.87,y,5.83,.34,value,14.5,INK)
line(s,8.04,2.45,8.04,6.68,LINE,.7)
rect(s,8.42,2.98,4.26,3.69,GREEN,radius=True)
txt(s,8.77,3.29,3.56,.28,'로컬 시연에서 확인한 결과',12,WHITE,True)
txt(s,8.77,3.84,3.56,.61,'CVE 3건 → 0건',29,WHITE,True)
txt(s,8.77,4.67,3.54,.29,'지정 Python 데모 환경의 재분석 결과',11.3,WHITE)
line(s,8.77,5.21,12.32,5.21,WHITE,.7)
txt(s,8.77,5.55,3.56,.79,'라이브러리 업데이트 전후 비교\n담당자 조치 기록 저장·재조회\n비교 PDF·JSON 다운로드 확인',13.2,WHITE,spacing=1.26)
footer(s,4)
s.notes_slide.notes_text_frame.text = '''적용은 등록한 고객사 자산과 분석 범위를 정하는 것에서 시작합니다. SSH로 읽을 서버 경로나 분석할 소스 ZIP을 지정하고, 생성한 결과를 해당 자산과 연결합니다. 기존 SPDX 또는 CycloneDX SBOM 반입도 지원합니다.
Syft가 구성요소를 식별해 SPDX 2.3 문서를 만들고, Grype가 알려진 취약점을 대조합니다. OSV는 선택적인 취약점 정보 보강에 사용합니다. 분석 도구 결과를 고객사 자산, 지원 종료 일정, 담당자 조치와 연결하는 관리 기능이 EOLWatch의 구현 범위입니다.
React 웹, FastAPI API와 worker, PostgreSQL을 Docker Compose로 구성합니다. 예약 분석과 DB·참조 ZIP의 백업·임시 복원 검증 기능도 구현했습니다. 대상 프로그램을 임의 실행하거나 자동으로 패치하지 않습니다.
검증 근거: 관리 VM 1대와 대상 VM 2대의 로컬 환경에서 웹 요청·SSH 수집·분석 결과 저장을 확인했습니다. 2026-09-15 지정 Python 가상환경의 Jinja2를 3.1.4에서 3.1.6으로 업데이트한 뒤, 같은 범위의 분석 #3과 #4에서 CVE 3건에서 0건으로 바뀌었습니다. 담당자 수동 조치 기록(action #1)의 저장·재조회와 비교 PDF·JSON 다운로드도 확인했습니다.
이 수치는 해당 데모 환경과 분석 시점의 재분석 미검출 결과입니다. 고객사 운영 환경의 도입 실적이나 전체 시스템의 안전성 보장을 의미하지 않습니다. 공개 EOL 개별 적용·분석 취소 등 추가 브라우저 흐름은 별도 검증 항목으로 남아 있으므로 전 기능 검증 완료라고 설명하지 않습니다.
발표 시에는 대응 대상 목록, 담당자별 작업목록, 조치 근거와 비교 보고서를 기업에 제공할 수 있는 결과물로 강조합니다.'''

# Ensure the inherited template adds no shadows or old master artwork.
for slide in prs.slides:
    for sh in slide.shapes:
        if sh.has_text_frame:
            for paragraph in sh.text_frame.paragraphs:
                for run in paragraph.runs:
                    if any('가' <= ch <= '힣' for ch in run.text):
                        run.font.name = KOREAN
        sppr=sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if sppr is not None:
            for tag in ['effectLst','effectDag']:
                for child in list(sppr.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag)):
                    sppr.remove(child)
            sppr.append(OxmlElement('a:effectLst'))
        for effect in sh._element.xpath('./p:style/a:effectRef'):
            effect.set('idx','0')
prs.core_properties.title='EOLWatch — 고객사 자산의 취약점·지원 종료·조치 관리'
prs.core_properties.subject='시스원 유지보수 업무를 위한 기능과 제공 결과'
prs.core_properties.author='김연동'
# Copy the finished editable slide content into a clean package, dropping unused
# media/layout relationships inherited from the supplied template.
clean = Presentation()
clean.slide_width, clean.slide_height = prs.slide_width, prs.slide_height
for source_slide in prs.slides:
    target = clean.slides.add_slide(clean.slide_layouts[6])
    target._element.replace(target._element.cSld, deepcopy(source_slide._element.cSld))
    target._element.set('showMasterSp','0')
    target.notes_slide.notes_text_frame.text = source_slide.notes_slide.notes_text_frame.text
for key in ['title','subject','author']:
    setattr(clean.core_properties,key,getattr(prs.core_properties,key))
prs = clean
prs.save(OUT)

check=Presentation(OUT)
assert len(check.slides)==4
allowed={'FAF9F9','4F785F','333333','A39D9D','C8C8C8'}
content=[]
for i,sl in enumerate(check.slides,1):
    assert set(sl._element.xpath('.//a:srgbClr/@val'))<=allowed
    for sh in sl.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=check.slide_width+100 and sh.top+sh.height<=check.slide_height+100,(i,sh.name)
    content.append({'slide':i,'text':[sh.text for sh in sl.shapes if sh.has_text_frame and sh.text],
                    'notes':sl.notes_slide.notes_text_frame.text})
all_text='\n'.join('\n'.join(s['text'])+s['notes'] for s in content)
assert not any(word in all_text for word in ['StockCast','NginX 방어','연동 검토','17개 테이블'])
assert hashlib.sha256(BASE.read_bytes()).hexdigest()==base_hash
(WORK/'content_review.json').write_text(json.dumps(content,ensure_ascii=False,indent=2))
(WORK/'validation.json').write_text(json.dumps({'source':str(BASE),'source_sha256':base_hash,
    'output':str(OUT),'slides':4,'palette':sorted(allowed),'source_unchanged':True},ensure_ascii=False,indent=2))
print(OUT)
print('Separate copy created; 4 slides rewritten for business value, source preserved.')
