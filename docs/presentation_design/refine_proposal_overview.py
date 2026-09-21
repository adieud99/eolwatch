"""Refine the project's role, restore its original four features, and plan validation."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from pptx import Presentation
import ast, hashlib, json, posixpath

HELPER = Path(__file__).resolve().with_name('reuse_base_template.py')
HERE = HELPER.parent
# Load pure XML utilities only; no dependency on deleted earlier deliverables.
module = ast.parse(HELPER.read_text())
selected = [node for node in module.body
            if (isinstance(node, ast.FunctionDef) and node.name in ('sha', 'xml', 'encode', 'replace_text'))
            or (isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'NS' for t in node.targets))]
exec(compile(ast.Module(body=selected, type_ignores=[]), str(HELPER), 'exec'), globals())
SOURCE = Path(json.loads((HERE / 'preproject_proposal/delivered_files.json').read_text())[0])
WORK = HERE / 'proposal_refined'
WORK.mkdir(exist_ok=True)
OUT = WORK / 'EOLWatch_프로젝트제안요약서_김연동_개요수정.pptx'
source_hash = sha(SOURCE)
with ZipFile(SOURCE) as z:
    original = {n: z.read(n) for n in z.namelist()}
parts = dict(original)

TEXT = {
    2: {
        4: '프로젝트 개요',
        5: '소프트웨어 현황과 조치 업무를 연결하는 관리 서비스',
        21: 'EOLWatch 소개',
        14: 'EOLWatch는 고객사 자산의 소프트웨어 구성·취약점·지원 종료와 조치 이력을 함께 관리하는 웹 서비스입니다.\n분석 결과를 실제 자산과 담당자 업무에 연결해, 대상 파악부터 조치·보고까지 지원하도록 개발하고자 합니다.',
        22: '프로젝트의 역할',
        18: '자산 현황 정리',
        19: '대응 판단 지원',
        20: '처리 근거 관리',
        15: '고객사·서버별 구성요소와\n취약점·지원 종료 정보를\n한곳에서 확인하도록 지원',
        16: '영향받는 자산과 변경 필요\n대상을 파악해 담당자가\n조치 순서를 정하도록 지원',
        17: '작업 이력과 재분석 결과를\n연결해 담당자 인계 및\n고객 보고 근거로 활용',
    },
    3: {
        4: '주요 기능',
        5: 'SBOM 생성부터 영향 자산 확인과 조치 관리까지',
        22: 'SBOM 생성',
        24: '취약점 분석',
        23: '자산 연결',
        25: '조치 관리',
        18: '설치된 구성요소·버전을\n식별해 SPDX 2.3으로 저장',
        20: '알려진 취약점과 대조해\nCVE와 수정 버전 확인',
        19: '결과를 고객사·사이트·자산에\n연결해 영향받는 서버 표시',
        21: '담당자·기한·상태·이력을\n기록하고 재분석 결과 비교',
    },
    4: {
        10: '검증 계획과 산출물',
        11: '시범 환경에서 기능을 확인하고, 검토 가능한 결과물 제공',
        20: '검증 계획',
        21: '예상 산출물',
        12: '환경 구성: 관리 서버와 점검 대상 서버로 시범 환경 준비\n기본 확인: 구성요소 수집 결과와 자산 연결의 정확성 점검\n변경 확인: 업데이트 전후 비교와 조치 이력·보고서 확인',
        13: '웹 시제품: 자산·분석 결과·담당자 조치를 확인하는 관리 화면\n보고 자료: 구성요소 목록과 조치 전후 비교 보고서\n운영 자료: 소스 코드·설치 구성·사용 및 백업·복구 안내',
    },
}

NOTES = {
    2: '''먼저 프로젝트가 무엇인지 정의하고, 유지관리 업무에서 맡을 역할을 연결해 설명합니다. EOLWatch는 고객사 자산의 소프트웨어 구성, 알려진 취약점, 지원 종료와 담당자 조치 이력을 함께 관리하도록 개발하고자 하는 웹 서비스입니다.
프로젝트의 역할은 세 가지입니다. 첫째, 고객사·서버별 구성요소 및 위험 정보를 한곳에서 확인하게 하는 현황 정리입니다. 둘째, 영향받는 자산과 교체·업그레이드 검토 대상을 확인해 담당자가 대응 순서를 정하도록 돕는 판단 지원입니다. 셋째, 작업 이력과 동일 범위 재분석 결과를 연결해 담당자 인계와 고객 보고에 쓸 근거를 관리하는 것입니다.
다음 장에서는 이 역할을 구현할 네 가지 기능인 SBOM 생성, 취약점 분석, 자산 연결, 조치 관리를 설명합니다. 지원 종료 일정은 자산·제품 정보와 함께 관리하는 방향으로 설계하며, 공식 출처와 확인 시점을 남길 계획입니다.
시스원의 현재 내부 시스템이나 업무의 부족을 단정하지 않으며, 실제 기업 적용이나 시간·비용 절감 성과를 주장하지 않습니다. 이 장의 역할은 개발 목표입니다.''',
    3: '''원래 요약서에서 정한 네 가지 핵심 기능을 유지합니다. 네 기능 모두 이번 프로젝트에서 개발할 범위로 제안합니다.
SBOM 생성: 지정된 서버 경로나 프로젝트에서 소프트웨어 구성요소와 버전을 식별해 SPDX 2.3 형식으로 저장할 계획입니다. SBOM은 소프트웨어 구성요소 명세서입니다.
취약점 분석: 구성요소 정보를 알려진 취약점 데이터와 대조해 CVE와 수정 버전을 확인하도록 개발할 계획입니다. 구성요소 식별과 취약점 대조에는 Syft·Grype 등 공개 도구를 활용하고, 분석 결과를 업무에 연결하는 관리 기능을 중심으로 개발합니다.
자산 연결: 분석 결과를 고객사·사이트·자산에 연결해 영향받는 서버를 확인하도록 할 계획입니다. 제품과 지원 종료 일정도 연결 자산을 기준으로 확인하는 방향으로 설계합니다.
조치 관리: 담당자·기한·상태와 처리 이력을 기록하고, 같은 자산의 같은 범위를 재분석해 조치 전후 변화를 비교하도록 개발할 계획입니다. 실제 업데이트와 조치 완료 판단은 담당자가 수행하도록 설계합니다.
이 장은 구현 완료 보고가 아니며, 개발할 핵심 기능 네 가지를 설명하는 장입니다.''',
    4: '''마지막 장은 앞서 제시한 네 기능이 제대로 연결되는지 어떻게 확인하고, 회사가 어떤 결과물을 검토할 수 있을지 보여줍니다. 기술 이름을 나열하기보다 검증 계획과 예상 산출물을 중심으로 설명합니다.
시범 환경에는 관리 서버와 점검 대상 서버를 준비할 계획입니다. 구체적인 대수·접근 권한·분석 범위는 요구사항을 정할 때 확인하며, 실제 고객사 운영 환경 접근이 확정됐다는 의미는 아닙니다.
기본 확인에서는 구성요소 수집 결과를 대상 환경의 실제 이름·버전과 대조하고, 분석 결과가 올바른 고객사·사이트·자산에 연결되는지 확인할 계획입니다.
변경 확인에서는 시범 환경의 소프트웨어를 업데이트한 뒤 같은 범위를 다시 분석할 계획입니다. 전후 결과가 구분되는지, 담당자·기한·작업 내용이 저장되는지, 보고서에 필요한 근거가 포함되는지 확인하고 보완하는 것이 목표입니다. 재분석에서 검출되지 않았다는 이유만으로 자동 완료나 안전성을 보장하지 않습니다.
예상 산출물은 자산·분석 결과·조치 작업을 확인하는 웹 시제품, 구성요소 목록과 조치 전후 비교 보고서, 소스 코드와 설치·사용·백업·복구 안내입니다. 이를 통해 기업 담당자가 기능과 자료의 적합성을 확인하고 향후 적용 여부를 검토할 수 있도록 할 계획입니다.
구현은 React·FastAPI·PostgreSQL과 Docker 기반 구성을 검토하며, 구성요소 식별·취약점 대조에는 공개 분석 도구를 활용할 계획입니다. 현재 검증 결과나 기업 도입 성과를 제시하는 장이 아닙니다.''',
}

for number, replacements in TEXT.items():
    path = f'ppt/slides/slide{number}.xml'
    root = xml(parts[path])
    for shape_id, value in replacements.items():
        replace_text(root, shape_id, value)
    parts[path] = encode(root)

for number, value in NOTES.items():
    rels = xml(parts[f'ppt/slides/_rels/slide{number}.xml.rels'])
    rel = next(r for r in rels if r.get('Type').endswith('/notesSlide'))
    path = posixpath.normpath(posixpath.join('ppt/slides', rel.get('Target')))
    root = xml(parts[path])
    body = root.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:txBody', namespaces=NS)[0]
    for paragraph in list(body.findall('a:p', NS)):
        body.remove(paragraph)
    for line in value.split('\n'):
        paragraph = E.SubElement(body, f'{{{NS["a"]}}}p')
        run = E.SubElement(paragraph, f'{{{NS["a"]}}}r')
        E.SubElement(run, f'{{{NS["a"]}}}t').text = line
    parts[path] = encode(root)

core = xml(parts['docProps/core.xml'])
core.find('{http://purl.org/dc/elements/1.1/}subject').text = '프로젝트 개요와 역할, 핵심 기능 네 가지, 검증 계획과 예상 산출물'
parts['docProps/core.xml'] = encode(core)

# Every shape, group, font, palette, and relationship remains as in the copied base.
for number in (2, 3, 4):
    path = f'ppt/slides/slide{number}.xml'
    before, after = xml(original[path]), xml(parts[path])
    for root in (before, after):
        for body in root.xpath('.//p:txBody', namespaces=NS):
            body.getparent().remove(body)
    assert E.tostring(before) == E.tostring(after), number
assert parts['ppt/slides/slide1.xml'] == original['ppt/slides/slide1.xml']

with ZipFile(OUT, 'w', ZIP_DEFLATED) as z:
    for name, data in parts.items():
        z.writestr(name, data)
check = Presentation(OUT)
assert len(check.slides) == 4
assert sha(SOURCE) == source_hash
content = [{'slide': n, 'text': [sh.text for sh in s.shapes if sh.has_text_frame and sh.text],
            'notes': s.notes_slide.notes_text_frame.text}
           for n, s in enumerate(check.slides, 1)]
feature_texts = content[2]['text']
assert all(feature in feature_texts for feature in ('SBOM 생성', '취약점 분석', '자산 연결', '조치 관리'))
(WORK / 'content_review.json').write_text(json.dumps(content, ensure_ascii=False, indent=2))
(WORK / 'validation.json').write_text(json.dumps({
    'source': str(SOURCE), 'source_sha256': source_hash, 'output': str(OUT),
    'slides_including_cover': 4, 'source_unchanged': True,
    'cover_unchanged': True, 'all_shape_geometry_preserved': True,
    'original_four_features_restored': True, 'purpose': 'pre-project proposal',
}, ensure_ascii=False, indent=2))
print(OUT)
print('Exactly 4 slides; original cover/layout preserved; original four features restored.')
